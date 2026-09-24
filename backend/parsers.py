# -*- coding: utf-8 -*-
"""
墨阁 · 文档解析模块
========================================================
作用：把 txt / docx / pdf / xlsx 读成纯文本，供「素材入库」使用。

三条设计原则（都跟隐私有关）：
1. 只用本地库解析，不调用任何云端服务 —— 素材内容不出你的电脑。
2. 解析失败不抛异常，而是返回带原因的字典，避免一个坏文件搞崩整个导入。
3. 输出结构统一，方便直接写进数据库。

返回的字典长这样：
    {
      "ok":     True / False,          # 是否成功
      "path":   文件绝对路径,
      "name":   文件名,
      "ext":    扩展名,               # 如 ".docx"
      "text":   提取出的纯文本,
      "chars":  正文字符数,
      "blocks": 段落/页数,            # docx 是段落数，pdf 是页数
      "note":   提示信息,             # 例如"含 12 张图片，未提取"
    }

单独试一个文件：
    python backend/parsers.py "D:\\素材\\分类\\打斗.docx"
"""

import os

# 支持的扩展名（对外展示用）
SUPPORTED_EXTS = (".txt", ".md", ".docx", ".pdf", ".xlsx", ".xlsm")

# 纯文本文件尝试的编码顺序（中文环境最常见的几种）
TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "big5", "utf-16")


# ----------------------------------------------------------------------
# 内部工具
# ----------------------------------------------------------------------

def _result(ok, path, text="", blocks=0, note="", ext=None):
    """统一构造返回值"""
    return {
        "ok": ok,
        "path": os.path.abspath(path),
        "name": os.path.basename(path),
        "ext": ext if ext is not None else os.path.splitext(path)[1].lower(),
        "text": text,
        "chars": len(text),
        "blocks": blocks,
        "note": note,
    }


def _tidy(text):
    """去掉过多空行，让正文更干净（不影响内容）"""
    lines = [ln.rstrip() for ln in text.splitlines()]
    out, blank = [], 0
    for ln in lines:
        if ln.strip():
            blank = 0
            out.append(ln)
        else:
            blank += 1
            if blank <= 1:          # 连续空行最多保留一个
                out.append("")
    return "\n".join(out).strip()


def _read_text_file(path):
    """读纯文本，自动试几种编码"""
    with open(path, "rb") as f:
        raw = f.read()
    for enc in TEXT_ENCODINGS:
        try:
            return raw.decode(enc), enc
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="ignore"), "utf-8（有字符读不出，已忽略）"


# ----------------------------------------------------------------------
# 各格式的解析
# ----------------------------------------------------------------------

def parse_txt(path):
    try:
        text, enc = _read_text_file(path)
    except Exception as e:
        return _result(False, path, note=f"读取失败：{e}")
    note = f"编码识别为 {enc}"
    if enc.startswith("utf-8（"):
        note += "；建议另存为 UTF-8 编码，避免个别字丢失"
    return _result(True, path, _tidy(text), note=note)


def _parse_docx_via_xml(path):
    """兜底方案：绕过 python-docx，直接读 word/document.xml 抽文字。

    什么时候会用到：文件里有坏掉的关系条目（例如图片被删过，
    留下 Target="../NULL"），python-docx 会直接报
    "There is no item named 'NULL' in the archive"。
    这种文件其实文字完好，只是结构有瑕疵，所以要有兜底。

    返回 (text, 段落数, 图片数)；失败返回 (None, 0, 0)
    """
    import html
    import re
    import zipfile

    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
            n_img = len([n for n in z.namelist() if n.startswith("word/media/")])
    except Exception:
        return None, 0, 0

    # 换行 / 制表符还原成字符
    xml = xml.replace("<w:br/>", "\n").replace("<w:br />", "\n")
    xml = xml.replace("<w:tab/>", "\t").replace("<w:tab />", "\t")

    paras = []
    for chunk in xml.split("</w:p>"):                 # 每段 <w:p> 是一段
        texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", chunk, flags=re.S)
        if texts:
            paras.append(html.unescape("".join(texts)))

    return "\n".join(paras), len(paras), n_img


def parse_docx(path):
    try:
        from docx import Document
    except ImportError:
        return _result(False, path, note="缺少 python-docx，先执行："
                                         "pip install python-docx")

    doc = None
    try:
        doc = Document(path)
    except Exception as e:
        # 常规路径失败 → 走兜底
        text, n_para, n_img = _parse_docx_via_xml(path)
        if text and text.strip():
            notes = [f"常规解析失败（{e}），已用兜底方式读出文字",
                     f"段落 {n_para} 个"]
            if n_img:
                notes.append(f"含 {n_img} 张图片（未提取）")
            return _result(True, path, _tidy(text), blocks=n_para,
                           note="；".join(notes))
        return _result(False, path,
                       note=f"打不开：{e}（若是老式 .doc，请先用 Word 另存为 .docx）")

    parts = []

    # 正文段落
    for p in doc.paragraphs:
        parts.append(p.text)

    # 表格里的文字（很多素材是表格形式记的，不能漏）
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = _tidy("\n".join(parts))

    # 提一下图片数量 —— 图片里的文字这个模块读不出来，先如实告知
    try:
        n_img = len(doc.inline_shapes)
    except Exception:
        n_img = 0

    notes = [f"段落 {len(doc.paragraphs)} 个"]
    if n_img:
        notes.append(f"含 {n_img} 张图片（图片内容未提取，只取了文字）")
    if not text:
        notes.append("没有提取到文字 —— 可能是纯图片文档")

    return _result(True, path, text, blocks=len(doc.paragraphs),
                   note="；".join(notes))


def parse_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        return _result(False, path, note="缺少 pdfplumber，先执行："
                                         "pip install pdfplumber")

    pages_text = []
    n_img = 0
    try:
        with pdfplumber.open(path) as pdf:
            n_pages = len(pdf.pages)
            for pg in pdf.pages:
                try:
                    pages_text.append(pg.extract_text() or "")
                except Exception:
                    pages_text.append("")
                try:
                    n_img += len(pg.images)          # 用来判断是不是扫描件
                except Exception:
                    pass
    except Exception as e:
        return _result(False, path, note=f"打不开：{e}")

    text = _tidy("\n".join(pages_text))
    notes = []

    if len(text) < 10:
        # 这是扫描件最常见的情况：PDF 里其实是一张张图片，没有文字层
        msg = "没有文字层，基本可以确定是扫描件/图片版 PDF —— 纯解析读不出字"
        if n_img:
            msg += f"（本文件内含 {n_img} 张图）"
        msg += "，需要用 OCR"
        notes.append(msg)
    else:
        notes.append(f"{n_pages} 页")

    # 有些 PDF 是"部分页扫描"，这里提示一下比例
    empty = sum(1 for t in pages_text if len(t.strip()) < 5)
    if 0 < empty < len(pages_text):
        notes.append(f"其中 {empty} 页没有文字（可能是图片页）")

    return _result(True, path, text, blocks=n_pages, note="；".join(notes))


def parse_xlsx(path):
    """读 Excel 表格。素材常常是表格形式记的，所以按行拼成文本。"""
    try:
        import openpyxl
    except ImportError:
        return _result(False, path, note="缺少 openpyxl，先执行："
                                         "pip install openpyxl")

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        return _result(False, path,
                       note=f"打不开：{e}（老式 .xls 请先另存为 .xlsx）")

    lines = []
    n_rows = 0
    try:
        for ws in wb.worksheets:
            lines.append(f"【工作表：{ws.title}】")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    lines.append(" | ".join(cells).rstrip(" |"))
                    n_rows += 1
    except Exception as e:
        return _result(False, path, note=f"读取中断：{e}")
    finally:
        try:
            wb.close()
        except Exception:
            pass

    text = _tidy("\n".join(lines))
    note = f"{len(wb.worksheets)} 个工作表、{n_rows} 行有内容"
    return _result(True, path, text, blocks=n_rows, note=note)


# ----------------------------------------------------------------------
# 对外入口
# ----------------------------------------------------------------------

PARSERS = {
    ".txt": parse_txt,
    ".md": parse_txt,
    ".docx": parse_docx,
    ".pdf": parse_pdf,
    ".xlsx": parse_xlsx,
    ".xlsm": parse_xlsx,
}


def parse_file(path):
    """解析一个文件，返回统一结构的字典。永远不抛异常。"""
    if not os.path.isfile(path):
        return _result(False, path, note="文件不存在")

    ext = os.path.splitext(path)[1].lower()

    parser = PARSERS.get(ext)
    if parser is None:
        return _result(False, path,
                       note=f"暂不支持的格式 {ext or '（无扩展名）'}，"
                            f"目前支持：{'、'.join(SUPPORTED_EXTS)}")
    return parser(path)


def scan_folder(folder, recursive=True):
    """扫一个文件夹里所有能解析的文件，返回路径列表（不解析内容）"""
    found = []
    if recursive:
        for root, _dirs, files in os.walk(folder):
            for fn in files:
                if os.path.splitext(fn)[1].lower() in PARSERS:
                    found.append(os.path.join(root, fn))
    else:
        for fn in os.listdir(folder):
            p = os.path.join(folder, fn)
            if os.path.isfile(p) and os.path.splitext(fn)[1].lower() in PARSERS:
                found.append(p)
    return sorted(found)


# ----------------------------------------------------------------------
# 命令行自测：python backend/parsers.py "某个文件或文件夹"
# ----------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    if len(sys.argv) < 2:
        print("用法：")
        print('    python backend/parsers.py "文件或文件夹"           只报告字数等信息')
        print('    python backend/parsers.py "单个文件" preview       额外打印正文前 200 字')
        raise SystemExit

    target = sys.argv[1]
    want_preview = len(sys.argv) > 2 and sys.argv[2].lower() in ("preview", "-p", "--preview")

    if os.path.isdir(target):
        files = scan_folder(target)
        print(f"文件夹里有 {len(files)} 个可解析文件：\n")
        for f in files:
            print("  -", f)
        print("\n" + "=" * 58)
        for f in files:
            r = parse_file(f)
            state = "OK  " if r["ok"] else "失败"
            print(f"[{state}] {r['name']}  {r['chars']} 字  {r['note']}")
    else:
        r = parse_file(target)
        print("=" * 58)
        print("文件：", r["name"])
        print("格式：", r["ext"])
        print("成功：", r["ok"])
        print("字数：", r["chars"])
        print("段/页：", r["blocks"])
        print("说明：", r["note"])
        print("=" * 58)
        if want_preview:
            print("正文前 200 字预览：")
            print(r["text"][:200])
        else:
            print("（正文未打印。想看前 200 字，命令末尾加一个 preview）")
